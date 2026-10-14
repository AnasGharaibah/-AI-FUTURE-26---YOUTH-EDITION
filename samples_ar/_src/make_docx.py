"""Builds national_id_service.docx (fictional sample). Run: uv run python samples_ar/_src/make_docx.py"""
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement

OUT = Path(__file__).resolve().parent.parent / "national_id_service.docx"
d = Document()


def rtl(p):
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    pPr = p._p.get_or_add_pPr()
    pPr.append(OxmlElement("w:bidi"))
    return p


def para(text, style=None):
    return rtl(d.add_paragraph(text, style=style))


d.add_heading("دليل خدمة إصدار بطاقة الهوية الوطنية — نموذج تجريبي", 1)
para("وثيقة تجريبية خيالية لأغراض الاختبار فقط.")
d.add_heading("الشروط", 2)
para("يحق لكل مواطن أتم 16 سنة من عمره الحصول على بطاقة الهوية الوطنية. تكون البطاقة صالحة لمدة 10 سنوات.")
d.add_heading("الوثائق المطلوبة", 2)
para("أولاً: لإصدار بطاقة الهوية لأول مرة:")
for t in ("شهادة الميلاد.", "دفتر العائلة.", "صورة شخصية واحدة حديثة."):
    para(t, "List Number")
para("ثانياً: لبدل فاقد أو تالف:")
for t in ("تقرير فقدان من الجهة الأمنية في حالة الفقدان.", "البطاقة التالفة في حالة التلف.", "صورة شخصية واحدة حديثة."):
    para(t, "List Number 2")
d.add_heading("الرسوم", 2)
rows = [("نوع الطلب", "الرسوم (دينار)"), ("إصدار لأول مرة", "10"), ("تجديد", "10"), ("بدل فاقد أو تالف", "20")]
tb = d.add_table(rows=len(rows), cols=2)
tb.style = "Table Grid"
for i, r in enumerate(rows):
    for j, v in enumerate(r):
        tb.cell(i, j).text = v
d.add_heading("الخطوات", 2)
for t in ("مراجعة أقرب مركز من مراكز الأحوال المدنية.", "تقديم الوثائق المطلوبة.", "دفع الرسوم في المركز.",
          "أخذ البصمة والصورة.", "استلام البطاقة بعد انتهاء المدة."):
    para(t, "List Number 3")
d.add_heading("مدة الإنجاز", 2)
para("3 أيام عمل لجميع أنواع الطلبات.")
d.add_heading("مكان التقديم", 2)
para("هذه الخدمة لا تقدم إلكترونياً، ويتم التقديم حصراً في مراكز الأحوال المدنية.")
d.save(OUT)
print("wrote", OUT)
